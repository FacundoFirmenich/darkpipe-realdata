#!/usr/bin/env python3
"""IRP-27 Gate 009 discovery.

Acquire ENARGAS' official Gas Entregado workbook (customer/volume weights source),
archive and inspect it, and discover official INDEC / Datos Argentina routes that
could contain a gas-specific price outcome. This script does not estimate effects.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse

import openpyxl
import requests
from bs4 import BeautifulSoup

GED_URL = "https://www.enargas.gob.ar/secciones/transporte-y-distribucion/datos-estadisticos/GED/GED.xlsx"
INDEC_IPC_PAGE = "https://www.indec.gob.ar/indec/web/Nivel4-Tema-3-5-31"
INDEC_SITEMAPS = [
    "https://www.indec.gob.ar/sitemap.xml",
    "https://www.indec.gob.ar/indec/web/sitemap.xml",
]
DISCOVERY_ENDPOINTS = [
    "https://apis.datos.gob.ar/series/api/search/?q=gas&limit=500",
    "https://apis.datos.gob.ar/series/api/search/?q=ipc%20gas&limit=500",
    "https://apis.datos.gob.ar/series/api/search/?q=indice%20precios%20consumidor%20gas&limit=500",
    "https://datos.gob.ar/api/3/action/package_search?q=gas%20ipc&rows=100",
    "https://datos.gob.ar/api/3/action/package_search?q=indice%20precios%20consumidor&rows=100",
]


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(path: Path, rows: Iterable[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def session() -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "User-Agent": "IRP27-Gate009/0.1 (+reproducible-public-data-discovery)",
        "Accept": "*/*",
    })
    return s


def download(s: requests.Session, url: str, path: Path, attempts: int = 3) -> dict[str, Any]:
    path.parent.mkdir(parents=True, exist_ok=True)
    last: str | None = None
    for attempt in range(1, attempts + 1):
        try:
            r = s.get(url, timeout=(20, 180), allow_redirects=True)
            r.raise_for_status()
            path.write_bytes(r.content)
            return {
                "requested_url": url,
                "final_url": r.url,
                "http_status": r.status_code,
                "content_type": r.headers.get("content-type"),
                "content_length_header": r.headers.get("content-length"),
                "etag": r.headers.get("etag"),
                "last_modified": r.headers.get("last-modified"),
                "bytes": len(r.content),
                "sha256": sha256_bytes(r.content),
                "retrieved_at_utc": utcnow(),
                "attempt": attempt,
                "error": None,
            }
        except Exception as exc:
            last = f"{type(exc).__name__}: {exc}"
            time.sleep(attempt * 2)
    return {
        "requested_url": url,
        "final_url": None,
        "http_status": None,
        "content_type": None,
        "bytes": 0,
        "sha256": None,
        "retrieved_at_utc": utcnow(),
        "attempt": attempts,
        "error": last,
    }


def fetch_text(s: requests.Session, url: str) -> tuple[dict[str, Any], str]:
    try:
        r = s.get(url, timeout=(20, 120), allow_redirects=True)
        meta = {
            "requested_url": url,
            "final_url": r.url,
            "http_status": r.status_code,
            "content_type": r.headers.get("content-type"),
            "bytes": len(r.content),
            "sha256": sha256_bytes(r.content),
            "retrieved_at_utc": utcnow(),
            "error": None,
        }
        r.raise_for_status()
        r.encoding = r.encoding or r.apparent_encoding or "utf-8"
        return meta, r.text
    except Exception as exc:
        return {
            "requested_url": url,
            "final_url": None,
            "http_status": None,
            "content_type": None,
            "bytes": 0,
            "sha256": None,
            "retrieved_at_utc": utcnow(),
            "error": f"{type(exc).__name__}: {exc}",
        }, ""


def inspect_ged(path: Path, out: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    row_samples: list[dict[str, Any]] = []
    nonempty_dump: list[dict[str, Any]] = []
    keywords = (
        "distrib", "subzona", "provincia", "usuario", "clientes", "volumen",
        "entregado", "mes", "año", "residencial", "domiciliarios", "servicio",
    )
    for ws in wb.worksheets:
        nonempty_rows = 0
        max_nonempty_col = 0
        header_candidates: list[dict[str, Any]] = []
        for row in ws.iter_rows():
            vals = [c.value for c in row]
            nonempty = [(i + 1, v) for i, v in enumerate(vals) if v not in (None, "")]
            if not nonempty:
                continue
            nonempty_rows += 1
            max_nonempty_col = max(max_nonempty_col, max(i for i, _ in nonempty))
            text = " | ".join(re.sub(r"\s+", " ", str(v)).strip() for _, v in nonempty)
            lower = text.casefold()
            rec = {
                "sheet": ws.title,
                "row": row[0].row if row else None,
                "nonempty_cells": len(nonempty),
                "first_col": nonempty[0][0],
                "last_col": nonempty[-1][0],
                "row_text": text,
            }
            if len(nonempty_dump) < 300000:
                nonempty_dump.append(rec)
            if any(k in lower for k in keywords):
                header_candidates.append(rec)
            if len(row_samples) < 250 and nonempty_rows <= 50:
                row_samples.append(rec)
        sheets.append({
            "sheet": ws.title,
            "reported_max_row": ws.max_row,
            "reported_max_column": ws.max_column,
            "nonempty_rows": nonempty_rows,
            "max_nonempty_column": max_nonempty_col,
            "header_candidate_count": len(header_candidates),
            "header_candidates": header_candidates[:100],
        })

    write_csv(
        out / "ged_nonempty_rows.csv",
        nonempty_dump,
        ["sheet", "row", "nonempty_cells", "first_col", "last_col", "row_text"],
    )
    write_csv(
        out / "ged_row_samples.csv",
        row_samples,
        ["sheet", "row", "nonempty_cells", "first_col", "last_col", "row_text"],
    )
    meta = {
        "filename": path.name,
        "sha256": sha256_file(path),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": sheets,
        "effect_estimation_authorized": False,
    }
    write_json(out / "ged_workbook_manifest.json", meta)
    return meta


def extract_links(base_url: str, html: str) -> list[dict[str, Any]]:
    soup = BeautifulSoup(html, "lxml")
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for tag in soup.find_all(["a", "link", "script", "form"]):
        attr = "href" if tag.name in {"a", "link"} else "src" if tag.name == "script" else "action"
        raw = tag.get(attr)
        if not raw:
            continue
        url = urljoin(base_url, raw)
        text = re.sub(r"\s+", " ", tag.get_text(" ", strip=True))
        key = (url, tag.name)
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "tag": tag.name,
            "attribute": attr,
            "raw": raw,
            "url": url,
            "text": text,
            "suffix": Path(urlparse(url).path).suffix.lower(),
            "contains_ipc": "ipc" in (url + " " + text).casefold(),
            "contains_gas": "gas" in (url + " " + text).casefold(),
        })
    # Also catch URLs embedded in JavaScript/JSON.
    for match in re.finditer(r"(?:https?://[^\s\"'<>]+|/[^\s\"'<>]+\.(?:xlsx?|csv|zip|json|pdf))", html, re.I):
        raw = match.group(0).rstrip(")],;\\")
        url = urljoin(base_url, raw)
        key = (url, "embedded")
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "tag": "embedded",
            "attribute": "regex",
            "raw": raw,
            "url": url,
            "text": "",
            "suffix": Path(urlparse(url).path).suffix.lower(),
            "contains_ipc": "ipc" in url.casefold(),
            "contains_gas": "gas" in url.casefold(),
        })
    return rows


def discover_indec(s: requests.Session, out: Path) -> dict[str, Any]:
    indec_dir = out / "indec_discovery"
    indec_dir.mkdir(parents=True, exist_ok=True)
    meta, html = fetch_text(s, INDEC_IPC_PAGE)
    (indec_dir / "ipc_page.html").write_text(html, encoding="utf-8")
    links = extract_links(meta.get("final_url") or INDEC_IPC_PAGE, html)
    write_csv(
        indec_dir / "ipc_page_links.csv",
        links,
        ["tag", "attribute", "raw", "url", "text", "suffix", "contains_ipc", "contains_gas"],
    )

    script_results: list[dict[str, Any]] = []
    script_links: list[dict[str, Any]] = []
    for idx, row in enumerate([r for r in links if r["tag"] == "script"][:40], 1):
        smeta, text = fetch_text(s, row["url"])
        smeta["source_url"] = row["url"]
        script_results.append(smeta)
        if text:
            (indec_dir / "scripts" / f"script_{idx:02d}.js").parent.mkdir(parents=True, exist_ok=True)
            (indec_dir / "scripts" / f"script_{idx:02d}.js").write_text(text, encoding="utf-8")
            script_links.extend(extract_links(row["url"], text))
    write_json(indec_dir / "script_fetch_registry.json", script_results)
    write_csv(
        indec_dir / "script_discovered_links.csv",
        script_links,
        ["tag", "attribute", "raw", "url", "text", "suffix", "contains_ipc", "contains_gas"],
    )

    sitemap_results: list[dict[str, Any]] = []
    sitemap_hits: list[str] = []
    for idx, url in enumerate(INDEC_SITEMAPS, 1):
        smeta, text = fetch_text(s, url)
        sitemap_results.append(smeta)
        if text:
            (indec_dir / f"sitemap_{idx}.xml").write_text(text, encoding="utf-8")
            for loc in re.findall(r"<loc>(.*?)</loc>", text, re.I | re.S):
                if "ipc" in loc.casefold() or "precio" in loc.casefold() or "gas" in loc.casefold():
                    sitemap_hits.append(loc.strip())
    write_json(indec_dir / "sitemap_registry.json", sitemap_results)
    (indec_dir / "sitemap_candidate_urls.txt").write_text("\n".join(sorted(set(sitemap_hits))) + "\n", encoding="utf-8")

    return {
        "ipc_page": meta,
        "link_count": len(links),
        "file_link_count": sum(1 for r in links if r["suffix"] in {".xls", ".xlsx", ".csv", ".zip", ".json", ".pdf"}),
        "script_count_fetched": len(script_results),
        "script_discovered_link_count": len(script_links),
        "sitemap_candidate_count": len(set(sitemap_hits)),
    }


def discover_apis(s: requests.Session, out: Path) -> list[dict[str, Any]]:
    api_dir = out / "api_discovery"
    api_dir.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    for idx, url in enumerate(DISCOVERY_ENDPOINTS, 1):
        meta, text = fetch_text(s, url)
        rec = dict(meta)
        rec["query_url"] = url
        if text:
            (api_dir / f"endpoint_{idx:02d}.txt").write_text(text, encoding="utf-8")
            rec["text_mentions_gas"] = text.casefold().count("gas")
            rec["text_mentions_ipc"] = text.casefold().count("ipc")
            try:
                parsed = json.loads(text)
                write_json(api_dir / f"endpoint_{idx:02d}.json", parsed)
                rec["valid_json"] = True
                if isinstance(parsed, dict):
                    rec["top_level_keys"] = sorted(parsed.keys())
            except Exception as exc:
                rec["valid_json"] = False
                rec["json_error"] = f"{type(exc).__name__}: {exc}"
        results.append(rec)
    write_json(api_dir / "api_query_registry.json", results)
    return results


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", required=True, type=Path)
    args = ap.parse_args()
    root: Path = args.output
    raw = root / "raw"
    parsed = root / "parsed"
    raw.mkdir(parents=True, exist_ok=True)
    parsed.mkdir(parents=True, exist_ok=True)
    s = session()

    ged_path = raw / "GED.xlsx"
    ged_source = download(s, GED_URL, ged_path)
    ged_source["xlsx_zip_valid"] = bool(ged_path.exists() and zipfile.is_zipfile(ged_path))
    write_json(root / "ged_source_record.json", ged_source)
    ged_manifest: dict[str, Any] | None = None
    ged_error: str | None = None
    if ged_source["xlsx_zip_valid"]:
        try:
            ged_manifest = inspect_ged(ged_path, parsed)
        except Exception as exc:
            ged_error = f"{type(exc).__name__}: {exc}"

    indec = discover_indec(s, root)
    api_results = discover_apis(s, root)

    status = {
        "gate": "IRP27-GATE-OUTCOME-ACQUISITION-AND-PRETREND-FEASIBILITY-009",
        "phase": "SOURCE_DISCOVERY_AND_WEIGHTS_ACQUISITION",
        "generated_at_utc": utcnow(),
        "ged_acquired": ged_source.get("bytes", 0) > 0,
        "ged_xlsx_valid": ged_source.get("xlsx_zip_valid", False),
        "ged_parsed": ged_manifest is not None,
        "ged_sheet_count": ged_manifest.get("sheet_count") if ged_manifest else None,
        "ged_error": ged_error,
        "indec_discovery": indec,
        "api_queries": len(api_results),
        "api_query_successes": sum(1 for x in api_results if x.get("http_status") == 200),
        "effect_estimates_produced": 0,
        "causal_authority": "BLOCKED",
        "political_comparison": "BLOCKED",
        "automatic_publication": "BLOCKED",
    }
    write_json(root / "status.json", status)

    manifest = []
    for p in sorted(root.rglob("*")):
        if p.is_file() and p.name != "MANIFEST.sha256":
            manifest.append(f"{sha256_file(p)}  {p.relative_to(root).as_posix()}")
    (root / "MANIFEST.sha256").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))
    return 0 if status["ged_xlsx_valid"] and status["ged_parsed"] else 2


if __name__ == "__main__":
    sys.exit(main())
