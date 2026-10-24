#!/usr/bin/env python3
"""IRP-27 Gate 012: acquire and inspect prior ENARGAS access-to-information responses.

Purpose
-------
Materialize prior official responses that demonstrate or contain residential-user and
volume openings by tariff category, subzone, benefit and municipality. The program
archives the source page, target-entry HTML, response documents, PDF text, embedded
attachments and tabular files. It never treats historical category shares as current
October-2026 weights and produces no bill or causal estimate.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import mimetypes
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

import openpyxl
import requests
from bs4 import BeautifulSoup, Tag
from pypdf import PdfReader

PAGE_URL = "https://www.enargas.gob.ar/secciones/transparencia/solicitar-informacion.php"
TARGETS = {
    "AIP-2018-CATEGORY-LICENSEE": ["4125/18", "Usuarios Residenciales por Categoría"],
    "AIP-2020-AMBA-CONSUMPTION-BANDS": ["NO-2020-88456063", "consumo promedio por segmento tarifario"],
    "AIP-2021-JUJUY-TUCUMAN-CATEGORY-SUBZONE": ["NO-2021-24551200", "usuarios residenciales abiertos por categoría"],
    "AIP-2022-PBA-MUNICIPAL-CATEGORY": ["NO-2022-45174850", "discriminados por categoría y municipio"],
    "AIP-2022-DEMAND-BENEFIT-SUBZONE-A": ["NO-2022-109180478", "beneficiarios de tarifa social"],
    "AIP-2022-DEMAND-BENEFIT-SUBZONE-B": ["NO-2022-115038124", "régimen de zona fría"],
    "AIP-2022-MUNICIPAL-CATEGORY-VOLUME": ["NO-2022-135121203", "volumen facturado de los años 2018 a 2021"],
    "AIP-2023-LOW-CONSUMPTION": ["NO-2023-14236662", "consumos inferiores a 30m3"],
    "AIP-2024-AMBA-R-CATEGORIES": ["NO-2024-47527596", "categorías R1R21R22"],
    "AIP-2024-GEOSPATIAL-LAYERS": ["NO-2024-61552214", "capas de información georreferenciada"],
    "AIP-2025-DEMAND-SEGMENT": ["NO-2025-30871687", "demanda"],
}

SESSION = requests.Session()
SESSION.headers.update({
    "User-Agent": "IRP27-Gate012/2.0 public-information-archive-research",
    "Accept": "text/html,application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/zip,*/*",
})


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_name(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_")
    return value[:180] or "object"


def fetch(url: str) -> tuple[bytes | None, dict[str, Any]]:
    error = None
    for attempt in range(1, 4):
        try:
            response = SESSION.get(url, timeout=(20, 180), allow_redirects=True)
            metadata = {
                "requested_url": url,
                "final_url": response.url,
                "http_status": response.status_code,
                "content_type": response.headers.get("content-type"),
                "content_disposition": response.headers.get("content-disposition"),
                "content_length_header": response.headers.get("content-length"),
                "etag": response.headers.get("etag"),
                "last_modified": response.headers.get("last-modified"),
                "retrieved_at_utc": now(),
                "attempt": attempt,
            }
            response.raise_for_status()
            data = response.content
            metadata.update({"bytes": len(data), "sha256": sha256_bytes(data), "error": None})
            return data, metadata
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
            time.sleep(attempt * 2)
    return None, {
        "requested_url": url,
        "final_url": None,
        "http_status": None,
        "content_type": None,
        "retrieved_at_utc": now(),
        "attempt": 3,
        "bytes": 0,
        "sha256": None,
        "error": error,
    }


def closest_entry(anchor: Tag) -> Tag:
    """Choose the smallest useful ancestor containing a record heading and description."""
    candidates = []
    for parent in anchor.parents:
        if not isinstance(parent, Tag):
            continue
        text = parent.get_text(" ", strip=True)
        if 100 <= len(text) <= 8000 and ("Expediente:" in text or "Solicitud" in text):
            candidates.append(parent)
        if parent.name in {"body", "html"}:
            break
    return candidates[0] if candidates else anchor.parent


def infer_filename(url: str, metadata: dict[str, Any], index: int) -> str:
    disposition = metadata.get("content_disposition") or ""
    match = re.search(r"filename\*?=(?:UTF-8''|\")?([^\";]+)", disposition, re.I)
    if match:
        return safe_name(match.group(1).strip())
    path_name = Path(urlparse(metadata.get("final_url") or url).path).name
    if path_name and "." in path_name:
        return safe_name(path_name)
    content_type = (metadata.get("content_type") or "").split(";")[0]
    extension = mimetypes.guess_extension(content_type) or ".bin"
    return f"attachment_{index:02d}{extension}"


def inspect_xlsx(path: Path, out_dir: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        width = max((len(row) for row in rows), default=0)
        fields = [f"col_{i+1}" for i in range(width)] + ["source_row_number"]
        csv_path = out_dir / f"{safe_name(ws.title)}.csv"
        nonempty = 0
        with csv_path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=fields)
            writer.writeheader()
            for row_number, row in enumerate(rows, start=1):
                if not any(value not in (None, "") for value in row):
                    continue
                nonempty += 1
                record = {f"col_{i+1}": row[i] if i < len(row) else None for i in range(width)}
                record["source_row_number"] = row_number
                writer.writerow(record)
        sheets.append({
            "sheet": ws.title,
            "max_row_reported": ws.max_row,
            "max_column_reported": ws.max_column,
            "nonempty_rows": nonempty,
            "csv_path": csv_path.name,
        })
    return {"kind": "xlsx", "sheet_count": len(sheets), "sheets": sheets}


def inspect_pdf(path: Path, out_dir: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    pages = []
    text_parts = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append({"page": page_number, "text_chars": len(text)})
        text_parts.append(f"\n\n--- PAGE {page_number} ---\n{text}")
    text_path = out_dir / f"{path.stem}.txt"
    text_path.write_text("".join(text_parts), encoding="utf-8")

    embedded = []
    attachments_dir = out_dir / f"{path.stem}_embedded"
    attachments_dir.mkdir(exist_ok=True)
    try:
        attachments = reader.attachments
        for name, payloads in attachments.items():
            if not isinstance(payloads, list):
                payloads = [payloads]
            for index, payload in enumerate(payloads, start=1):
                child = attachments_dir / safe_name(f"{index:02d}_{name}")
                child.write_bytes(payload)
                embedded.append({
                    "name": name,
                    "path": str(child.relative_to(out_dir)),
                    "bytes": len(payload),
                    "sha256": sha256_bytes(payload),
                })
    except Exception as exc:
        embedded.append({"inspection_error": f"{type(exc).__name__}: {exc}"})
    return {"kind": "pdf", "page_count": len(pages), "pages": pages, "text_path": text_path.name, "embedded": embedded}


def inspect_object(path: Path, content_type: str | None, out_dir: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"} or "spreadsheetml" in (content_type or ""):
        return inspect_xlsx(path, out_dir)
    if suffix == ".pdf" or "pdf" in (content_type or ""):
        return inspect_pdf(path, out_dir)
    if suffix in {".zip"} or "zip" in (content_type or ""):
        import zipfile
        with zipfile.ZipFile(path) as zf:
            return {"kind": "zip", "members": zf.namelist()}
    return {"kind": "other"}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    root = args.output
    root.mkdir(parents=True, exist_ok=True)

    page_data, page_meta = fetch(PAGE_URL)
    if page_data is None:
        (root / "acquisition_status.json").write_text(json.dumps({"page": page_meta}, indent=2), encoding="utf-8")
        return 2
    (root / "solicitar-informacion.html").write_bytes(page_data)
    html = page_data.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    link_inventory = []
    for index, anchor in enumerate(soup.find_all("a", href=True), start=1):
        link_inventory.append({
            "index": index,
            "text": anchor.get_text(" ", strip=True),
            "href": urljoin(PAGE_URL, anchor["href"]),
        })
    with (root / "link_inventory.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=["index", "text", "href"])
        writer.writeheader(); writer.writerows(link_inventory)

    target_registry = []
    document_registry = []
    unmatched = []
    for target_id, needles in TARGETS.items():
        matching_anchors: list[Tag] = []
        for anchor in soup.find_all("a", href=True):
            anchor_text = anchor.get_text(" ", strip=True)
            parent_text = closest_entry(anchor).get_text(" ", strip=True)
            if any(needle.casefold() in anchor_text.casefold() or needle.casefold() in parent_text.casefold() for needle in needles):
                matching_anchors.append(anchor)

        unique_entries: dict[str, Tag] = {}
        for anchor in matching_anchors:
            entry = closest_entry(anchor)
            key = hashlib.sha256(entry.get_text(" ", strip=True).encode("utf-8")).hexdigest()
            unique_entries[key] = entry
        if not unique_entries:
            unmatched.append({"target_id": target_id, "needles": needles})
            continue

        target_dir = root / "targets" / target_id
        target_dir.mkdir(parents=True, exist_ok=True)
        entry_links: list[dict[str, str]] = []
        entry_texts = []
        for entry_index, entry in enumerate(unique_entries.values(), start=1):
            entry_html = str(entry)
            (target_dir / f"entry_{entry_index:02d}.html").write_text(entry_html, encoding="utf-8")
            entry_text = entry.get_text("\n", strip=True)
            (target_dir / f"entry_{entry_index:02d}.txt").write_text(entry_text, encoding="utf-8")
            entry_texts.append(entry_text)
            for anchor in entry.find_all("a", href=True):
                entry_links.append({
                    "text": anchor.get_text(" ", strip=True),
                    "url": urljoin(PAGE_URL, anchor["href"]),
                })

        # Deduplicate URLs while retaining all official record links in the target entry.
        dedup: dict[str, dict[str, str]] = {}
        for item in entry_links:
            dedup.setdefault(item["url"], item)
        entry_links = list(dedup.values())
        target_registry.append({
            "target_id": target_id,
            "needles": needles,
            "entries": len(unique_entries),
            "links": len(entry_links),
            "entry_text_sha256": sha256_bytes("\n\n".join(entry_texts).encode("utf-8")),
        })

        for document_index, item in enumerate(entry_links, start=1):
            data, metadata = fetch(item["url"])
            metadata.update({"target_id": target_id, "anchor_text": item["text"]})
            if data is None:
                document_registry.append(metadata)
                continue
            filename = infer_filename(item["url"], metadata, document_index)
            path = target_dir / filename
            path.write_bytes(data)
            metadata.update({"filename": filename, "local_sha256": sha256_file(path)})
            try:
                metadata["inspection"] = inspect_object(path, metadata.get("content_type"), target_dir)
            except Exception as exc:
                metadata["inspection"] = {"error": f"{type(exc).__name__}: {exc}"}
            document_registry.append(metadata)

    (root / "page_metadata.json").write_text(json.dumps(page_meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (root / "target_registry.json").write_text(json.dumps(target_registry, ensure_ascii=False, indent=2), encoding="utf-8")
    (root / "document_registry.json").write_text(json.dumps(document_registry, ensure_ascii=False, indent=2), encoding="utf-8")
    (root / "unmatched_targets.json").write_text(json.dumps(unmatched, ensure_ascii=False, indent=2), encoding="utf-8")

    status = {
        "gate": "IRP27-GATE-EXACT-WEIGHT-VECTOR-AND-DISTRIBUTIVE-INCIDENCE-012",
        "source_page_acquired": True,
        "targets_expected": len(TARGETS),
        "targets_matched": len(target_registry),
        "documents_attempted": len(document_registry),
        "documents_acquired": sum(int(row.get("bytes") or 0) > 0 for row in document_registry),
        "xlsx_objects": sum(row.get("inspection", {}).get("kind") == "xlsx" for row in document_registry),
        "pdf_objects": sum(row.get("inspection", {}).get("kind") == "pdf" for row in document_registry),
        "historical_weight_transport_authority": "BLOCKED_UNTIL_TIME_AND_SCOPE_MODEL",
        "october_2026_exact_weight_vector": "NOT_ACQUIRED_BY_THIS_ARCHIVE_PASS",
        "national_bill_average": "NOT_ESTIMATED",
        "automatic_publication": False,
        "generated_at_utc": now(),
    }
    (root / "acquisition_status.json").write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")

    manifest = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.name != "MANIFEST.sha256":
            manifest.append(f"{sha256_file(path)}  {path.relative_to(root).as_posix()}")
    (root / "MANIFEST.sha256").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))
    return 0 if status["targets_matched"] >= 7 else 2


if __name__ == "__main__":
    sys.exit(main())
