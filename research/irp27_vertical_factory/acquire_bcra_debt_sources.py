#!/usr/bin/env python3
"""Acquire and inspect official BCRA sources for IRP-27 Deuda Abierta.

The job archives raw PDFs/XLSX, hashes bytes, validates XLSX containers, and exports
lossless sheet CSVs. It does not enumerate individual debtors and does not estimate a
policy effect, portfolio purchase price, household payment burden, or fiscal cost.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen
import zipfile

import openpyxl
from pypdf import PdfReader

SOURCES = [
    {
        "source_id": "BCRA-BANKS-REPORT-2026-06",
        "kind": "pdf",
        "url": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/informe-bancos-2026-06.pdf",
    },
    {
        "source_id": "BCRA-BANKS-SERIES-2026-06",
        "kind": "xlsx",
        "url": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/informe-bancos-serie-2026-06.xlsx",
    },
    {
        "source_id": "BCRA-BANKS-ANNEX-2026-06",
        "kind": "xlsx",
        "url": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/informe-bancos-anexo.xlsx",
    },
    {
        "source_id": "BCRA-PNFC-REPORT-2026-06",
        "kind": "pdf",
        "url": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/informe-proveedores-no-financieros-credito-junio-2026.pdf",
    },
    {
        "source_id": "BCRA-PNFC-SERIES-2026-06",
        "kind": "xlsx",
        "url": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/series-informe-proveedores-no-financieros-credito-junio-2026.xlsx",
    },
    {
        "source_id": "BCRA-PNFC-ANNEX-2026-06",
        "kind": "xlsx",
        "url": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/anexo-estadistico-proveedores-no-financieros-credito-junio-2026.xlsx",
    },
]


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_") or "sheet"


def fetch(url: str, dest: Path) -> dict[str, Any]:
    err = None
    for attempt in range(1, 4):
        try:
            req = Request(url, headers={
                "User-Agent": "IRP27-Deuda-Abierta/0.2 public-data-acquisition",
                "Accept": "application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*",
            })
            with urlopen(req, timeout=180) as response:
                data = response.read()
                dest.write_bytes(data)
                return {
                    "requested_url": url,
                    "final_url": response.geturl(),
                    "http_status": getattr(response, "status", 200),
                    "content_type": response.headers.get("Content-Type"),
                    "etag": response.headers.get("ETag"),
                    "last_modified": response.headers.get("Last-Modified"),
                    "retrieved_at_utc": now(),
                    "attempt": attempt,
                    "bytes": len(data),
                    "sha256": sha256_bytes(data),
                    "error": None,
                }
        except Exception as exc:
            err = f"{type(exc).__name__}: {exc}"
            time.sleep(attempt * 2)
    return {
        "requested_url": url,
        "final_url": None,
        "http_status": None,
        "content_type": None,
        "etag": None,
        "last_modified": None,
        "retrieved_at_utc": now(),
        "attempt": 3,
        "bytes": 0,
        "sha256": None,
        "error": err,
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = list(rows[0].keys()) if rows else []
    with path.open("w", newline="", encoding="utf-8") as fh:
        if fields:
            writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)


def inspect_xlsx(path: Path, source_id: str, root: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    dump_root = root / "sheet_csv" / source_id
    dump_root.mkdir(parents=True, exist_ok=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        width = max((len(r) for r in rows), default=0)
        fields = [f"col_{i+1}" for i in range(width)]
        records = []
        nonempty = 0
        for row_index, row in enumerate(rows, start=1):
            if not any(v not in (None, "") for v in row):
                continue
            nonempty += 1
            record = {fields[i]: row[i] if i < len(row) else None for i in range(width)}
            record["source_row_number"] = row_index
            records.append(record)
        write_csv(dump_root / f"{safe(ws.title)}.csv", records, fields + ["source_row_number"])
        sheets.append({
            "sheet": ws.title,
            "max_row_reported": ws.max_row,
            "max_column_reported": ws.max_column,
            "nonempty_rows": nonempty,
            "csv_path": str((dump_root / f"{safe(ws.title)}.csv").relative_to(root)),
        })
    return {"valid_xlsx": zipfile.is_zipfile(path), "sheet_count": len(sheets), "sheets": sheets}


def inspect_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    extracted = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        extracted.append({"page": page_no, "text_chars": len(text), "text": text})
    return {"page_count": len(extracted), "pages": extracted}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    root = args.output
    raw = root / "raw"
    raw.mkdir(parents=True, exist_ok=True)

    registry = []
    inspections = []
    errors = []
    for spec in SOURCES:
        ext = spec["kind"]
        filename = spec["url"].split("/")[-1]
        path = raw / filename
        record = fetch(spec["url"], path)
        record.update({"source_id": spec["source_id"], "kind": ext, "filename": filename})
        if record["bytes"] == 0:
            errors.append({"source_id": spec["source_id"], "stage": "fetch", "error": record["error"]})
            registry.append(record)
            continue
        try:
            if ext == "xlsx":
                detail = inspect_xlsx(path, spec["source_id"], root)
            else:
                detail = inspect_pdf(path)
                text_path = root / "pdf_text" / f"{spec['source_id']}.json"
                text_path.parent.mkdir(parents=True, exist_ok=True)
                text_path.write_text(json.dumps(detail, ensure_ascii=False, indent=2), encoding="utf-8")
                detail = {"page_count": detail["page_count"], "text_json_path": str(text_path.relative_to(root))}
            inspections.append({"source_id": spec["source_id"], **detail})
        except Exception as exc:
            errors.append({"source_id": spec["source_id"], "stage": "inspect", "error": f"{type(exc).__name__}: {exc}"})
        registry.append(record)

    write_csv(root / "source_registry.csv", registry, [
        "source_id", "kind", "filename", "requested_url", "final_url", "http_status", "content_type",
        "bytes", "sha256", "etag", "last_modified", "retrieved_at_utc", "attempt", "error",
    ])
    (root / "source_registry.json").write_text(json.dumps(registry, ensure_ascii=False, indent=2), encoding="utf-8")
    (root / "inspection_registry.json").write_text(json.dumps(inspections, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(root / "errors.csv", errors, ["source_id", "stage", "error"])

    status = {
        "gate": "IRP27-VERTICAL-FACTORY-002-DEBT-SOURCE-ACQUISITION",
        "expected_sources": len(SOURCES),
        "acquired_sources": sum(r["bytes"] > 0 for r in registry),
        "xlsx_expected": sum(s["kind"] == "xlsx" for s in SOURCES),
        "xlsx_valid": sum(i.get("valid_xlsx") is True for i in inspections),
        "pdf_expected": sum(s["kind"] == "pdf" for s in SOURCES),
        "pdf_parsed": sum("page_count" in i for i in inspections if i["source_id"].endswith("REPORT-2026-06")),
        "individual_debtor_enumeration": False,
        "causal_effect_estimates": 0,
        "portfolio_purchase_price_estimates": 0,
        "fiscal_cost_estimates": 0,
        "automatic_publication": False,
        "errors": len(errors),
    }
    (root / "acquisition_status.json").write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")

    manifest = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.name != "MANIFEST.sha256":
            manifest.append(f"{sha256_file(path)}  {path.relative_to(root).as_posix()}")
    (root / "MANIFEST.sha256").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))
    return 0 if status["acquired_sources"] == status["expected_sources"] else 2


if __name__ == "__main__":
    sys.exit(main())
