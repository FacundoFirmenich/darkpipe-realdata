from pathlib import Path
import hashlib,json,urllib.request,zipfile,sys
from datetime import datetime,timezone
import openpyxl
URL="https://www.enargas.gob.ar/secciones/precios-y-tarifas/descargas/tarifas-gn-20261001.xlsx"
out=Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=True)
path=out/"tarifas-gn-20261001.xlsx"
with urllib.request.urlopen(urllib.request.Request(URL,headers={"User-Agent":"IRP27-Shadow/1.6"}),timeout=150) as response:
    body=response.read()
    path.write_bytes(body)
    status={"requested_url":URL,"final_url":response.geturl(),"status":response.status,"content_type":response.headers.get("Content-Type"),"retrieved_at":datetime.now(timezone.utc).isoformat(),"size_bytes":len(body),"sha256":hashlib.sha256(body).hexdigest(),"valid_xlsx":zipfile.is_zipfile(path)}
if status["valid_xlsx"]:
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    status["sheets"]=wb.sheetnames
    status["sheet_row_counts"]={s.title:s.max_row for s in wb}
status.update({"automatic_publication":False,"release_signed":False,"causal_authority":False})
(out/"status.json").write_text(json.dumps(status,indent=2,ensure_ascii=False))
print(json.dumps(status,indent=2,ensure_ascii=False))
if not status["valid_xlsx"]:raise SystemExit(2)
