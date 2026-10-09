#!/usr/bin/env python3
from __future__ import annotations
import csv, hashlib, json, sys, urllib.request, zipfile
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

URL = "https://www.enargas.gob.ar/secciones/precios-y-tarifas/descargas/tarifas-gn-20260901.xlsx"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("artifact")
    out.mkdir(parents=True, exist_ok=True)
    raw = out / "tarifas-gn-20260901.xlsx"
    req = urllib.request.Request(URL, headers={"User-Agent":"IRP27-Energy-Vertical/1.5"})
    with urllib.request.urlopen(req, timeout=180) as r:
        raw.write_bytes(r.read())
        meta = {
            "requested_url": URL,
            "final_url": r.geturl(),
            "status": getattr(r, "status", 200),
            "content_type": r.headers.get("Content-Type"),
            "etag": r.headers.get("ETag"),
            "last_modified": r.headers.get("Last-Modified"),
            "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        }
    meta.update({"bytes": raw.stat().st_size, "sha256": sha256(raw), "xlsx_valid": zipfile.is_zipfile(raw)})
    if not meta["xlsx_valid"]:
        raise SystemExit("Downloaded object is not a valid XLSX container")
    wb = openpyxl.load_workbook(raw, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(x or "").strip() for x in next(it)]
        except StopIteration:
            continue
        for i, values in enumerate(it, start=2):
            if not any(v not in (None, "") for v in values):
                continue
            record = {header[j]: values[j] if j < len(values) else None for j in range(len(header))}
            record["source_sheet"] = ws.title
            record["source_row_number"] = i
            rows.append(record)
    (out / "source_metadata.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "workbook_summary.json").write_text(json.dumps({
        "sheet_names": wb.sheetnames,
        "rows": len(rows),
        "headers": sorted({k for r in rows for k in r.keys()}),
        "operators": sorted({str(r.get("EMPRESA") or "") for r in rows if r.get("EMPRESA")}),
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    if rows:
        fields = list(rows[0].keys())
        with (out / "tarifas-gn-20260901.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            w.writeheader(); w.writerows(rows)
    manifest=[]
    for p in sorted(out.glob("*")):
        if p.is_file() and p.name != "MANIFEST.sha256":
            manifest.append(f"{sha256(p)}  {p.name}")
    (out / "MANIFEST.sha256").write_text("\n".join(manifest)+"\n", encoding="utf-8")
    print(json.dumps({**meta, "rows":len(rows), "sheets":len(wb.sheetnames)}, indent=2, ensure_ascii=False))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
