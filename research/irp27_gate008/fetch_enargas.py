#!/usr/bin/env python3
"""IRP-27 Gate 008: acquire, hash, validate and inspect official ENARGAS tariff workbooks.

This program deliberately does not estimate policy effects. It creates a reproducible
raw-source package, workbook schema inventory, sparse cell dump and operator-detection
ledger. Any unresolved download or parse problem is preserved explicitly.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import requests

BASE = "https://www.enargas.gob.ar/secciones/precios-y-tarifas/descargas"
SPECS = [
    ("2026-04-01", f"{BASE}/tarifas-gn-20260401.xlsx"),
    ("2026-05-01", f"{BASE}/tarifas-gn-20260501.xlsx"),
    ("2026-05-04", f"{BASE}/tarifas-gn-20260504.xlsx"),
    ("2026-06-01", f"{BASE}/tarifas-gn-20260601.xlsx"),
    ("2026-07-01", f"{BASE}/tarifas-gn-20260701.xlsx"),
    ("2026-08-01", f"{BASE}/tarifas-gn-20260801.xlsx"),
]

OPERATORS = {
    "NATURGY_NOA": ["naturgy noa", "gasnor"],
    "CAMUZZI_GAS_PAMPEANA": ["camuzzi gas pampeana", "pampeana"],
    "LITORAL_GAS": ["litoral gas"],
    "CAMUZZI_GAS_DEL_SUR": ["camuzzi gas del sur", "gas del sur"],
    "GAS_NEA": ["gas nea", "gnea"],
    "METROGAS": ["metrogas", "metro gas"],
    "DISTRIBUIDORA_GAS_CUYANA": ["distribuidora de gas cuyana", "gas cuyana", "ecogas cuyana"],
    "DISTRIBUIDORA_GAS_CENTRO": ["distribuidora de gas del centro", "gas del centro", "ecogas centro"],
    "NATURGY_BAN": ["naturgy ban", "gas natural ban"],
}

TARIFF_TERMS = (
    "cargo fijo", "cargo variable", "factura mínima", "factura minima",
    "tarifa", "categoría", "categoria", "residencial", "m3", "m³",
    "subzona", "usuario", "consumo", "bonificación", "bonificacion",
    "precio de gas", "pau", "transporte", "dda", "rqt",
)


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_") or "sheet"


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "User-Agent": "IRP27-Gate008/0.9 (+reproducible public-data acquisition)",
        "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*",
    })
    return s


def download(session: requests.Session, url: str, dest: Path) -> dict[str, Any]:
    last_error: str | None = None
    for attempt in range(1, 4):
        try:
            response = session.get(url, timeout=(20, 120), allow_redirects=True)
            record = {
                "requested_url": url,
                "final_url": response.url,
                "http_status": response.status_code,
                "content_type": response.headers.get("content-type"),
                "content_length_header": response.headers.get("content-length"),
                "etag": response.headers.get("etag"),
                "last_modified": response.headers.get("last-modified"),
                "attempt": attempt,
                "retrieved_at_utc": utcnow(),
            }
            response.raise_for_status()
            data = response.content
            dest.write_bytes(data)
            record.update({
                "bytes": len(data),
                "sha256": sha256_bytes(data),
                "xlsx_zip_valid": zipfile.is_zipfile(dest),
            })
            return record
        except Exception as exc:  # retain exact acquisition failure
            last_error = f"{type(exc).__name__}: {exc}"
            time.sleep(attempt * 2)
    return {
        "requested_url": url,
        "final_url": None,
        "http_status": None,
        "retrieved_at_utc": utcnow(),
        "error": last_error,
        "bytes": 0,
        "sha256": None,
        "xlsx_zip_valid": False,
    }


def inspect_workbook(path: Path, vintage: str, out_dir: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    workbook_rows: list[dict[str, Any]] = []
    candidate_rows: list[dict[str, Any]] = []
    operator_hits: dict[str, list[dict[str, Any]]] = {key: [] for key in OPERATORS}
    sheet_meta: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        sheet_meta.append({
            "sheet": ws.title,
            "max_row_reported": ws.max_row,
            "max_column_reported": ws.max_column,
        })
        for row in ws.iter_rows():
            values = [cell.value for cell in row]
            nonempty = [(idx + 1, value) for idx, value in enumerate(values) if value not in (None, "")]
            if not nonempty:
                continue
            joined = " | ".join(normalize_text(v) for _, v in nonempty)
            folded = joined.casefold()
            row_record = {
                "vintage": vintage,
                "sheet": ws.title,
                "row": row[0].row if row else None,
                "first_nonempty_column": nonempty[0][0],
                "nonempty_cells": len(nonempty),
                "row_text": joined,
            }
            workbook_rows.append(row_record)

            if any(term in folded for term in TARIFF_TERMS):
                candidate_rows.append(row_record)

            for operator, aliases in OPERATORS.items():
                if any(alias.casefold() in folded for alias in aliases):
                    operator_hits[operator].append({
                        "sheet": ws.title,
                        "row": row_record["row"],
                        "row_text": joined,
                    })

    dump_dir = out_dir / "sparse_dumps"
    write_csv(
        dump_dir / f"{vintage}_all_nonempty_rows.csv",
        workbook_rows,
        ["vintage", "sheet", "row", "first_nonempty_column", "nonempty_cells", "row_text"],
    )
    write_csv(
        dump_dir / f"{vintage}_tariff_candidate_rows.csv",
        candidate_rows,
        ["vintage", "sheet", "row", "first_nonempty_column", "nonempty_cells", "row_text"],
    )

    coverage = []
    for operator in OPERATORS:
        hits = operator_hits[operator]
        coverage.append({
            "vintage": vintage,
            "operator": operator,
            "detected": bool(hits),
            "hit_count": len(hits),
            "first_sheet": hits[0]["sheet"] if hits else None,
            "first_row": hits[0]["row"] if hits else None,
        })

    return {
        "vintage": vintage,
        "filename": path.name,
        "sha256": sha256_file(path),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheet_metadata": sheet_meta,
        "nonempty_row_count": len(workbook_rows),
        "tariff_candidate_row_count": len(candidate_rows),
        "operator_hits": operator_hits,
        "operator_coverage": coverage,
        "dose_reconstruction_authorized": False,
        "effect_estimation_authorized": False,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    root: Path = args.output
    raw_dir = root / "raw"
    parsed_dir = root / "parsed"
    raw_dir.mkdir(parents=True, exist_ok=True)
    parsed_dir.mkdir(parents=True, exist_ok=True)

    session = make_session()
    source_records: list[dict[str, Any]] = []
    workbook_manifests: list[dict[str, Any]] = []
    coverage_rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for vintage, url in SPECS:
        filename = f"tarifas-gn-{vintage.replace('-', '')}.xlsx"
        path = raw_dir / filename
        acquisition = download(session, url, path)
        acquisition.update({"vintage": vintage, "filename": filename})
        source_records.append(acquisition)

        if not acquisition.get("xlsx_zip_valid"):
            errors.append({
                "vintage": vintage,
                "stage": "download_or_xlsx_validation",
                "error": acquisition.get("error") or f"HTTP/content validation failed: {acquisition}",
            })
            continue

        try:
            manifest = inspect_workbook(path, vintage, root)
            workbook_manifests.append(manifest)
            coverage_rows.extend(manifest["operator_coverage"])
            (parsed_dir / f"{vintage}_workbook_manifest.json").write_text(
                json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
            )
        except Exception as exc:
            errors.append({
                "vintage": vintage,
                "stage": "openpyxl_inspection",
                "error": f"{type(exc).__name__}: {exc}",
            })

    write_csv(
        root / "raw_source_manifest.csv",
        source_records,
        [
            "vintage", "filename", "requested_url", "final_url", "http_status",
            "content_type", "content_length_header", "bytes", "sha256", "xlsx_zip_valid",
            "etag", "last_modified", "attempt", "retrieved_at_utc", "error",
        ],
    )
    (root / "raw_source_manifest.json").write_text(
        json.dumps(source_records, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    write_csv(
        root / "operator_vintage_detection.csv",
        coverage_rows,
        ["vintage", "operator", "detected", "hit_count", "first_sheet", "first_row"],
    )
    write_csv(root / "errors.csv", errors, ["vintage", "stage", "error"])

    acquired = sum(1 for row in source_records if row.get("bytes", 0) > 0)
    valid = sum(1 for row in source_records if row.get("xlsx_zip_valid"))
    parsed = len(workbook_manifests)
    detected_cells = sum(1 for row in coverage_rows if row.get("detected"))
    status = {
        "gate": "IRP27-GATE-RAW-SOURCE-MATERIALIZATION-AND-TARIFF-DOSE-RECONSTRUCTION-008",
        "generated_at_utc": utcnow(),
        "expected_workbooks": len(SPECS),
        "acquired_workbooks": acquired,
        "valid_xlsx_workbooks": valid,
        "parsed_workbooks": parsed,
        "expected_operator_vintage_cells": len(SPECS) * len(OPERATORS),
        "operator_vintage_cells_with_textual_detection": detected_cells,
        "dose_reconstruction_complete": False,
        "effect_estimates_produced": 0,
        "causal_authority": "BLOCKED",
        "automatic_publication": "BLOCKED",
        "errors": len(errors),
    }
    (root / "acquisition_status.json").write_text(
        json.dumps(status, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    manifest_lines: list[str] = []
    for p in sorted(root.rglob("*")):
        if p.is_file() and p.name != "MANIFEST.sha256":
            manifest_lines.append(f"{sha256_file(p)}  {p.relative_to(root).as_posix()}")
    (root / "MANIFEST.sha256").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")

    print(json.dumps(status, indent=2, ensure_ascii=False))
    return 0 if valid == len(SPECS) and parsed == len(SPECS) else 2


if __name__ == "__main__":
    sys.exit(main())
